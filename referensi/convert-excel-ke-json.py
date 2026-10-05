import json, re, datetime, os
from openpyxl import load_workbook
U='/root/.claude/uploads/794154df-f70d-56df-8850-7700d50936f6/'
BUL={'januari':1,'februari':2,'maret':3,'april':4,'mei':5,'juni':6,'juli':7,'agustus':8,'agust':8,'september':9,'oktober':10,'november':11,'desember':12,'jan':1,'feb':2,'mar':3,'apr':4,'jun':6,'jul':7,'agu':8,'agt':8,'sep':9,'okt':10,'nov':11,'des':12}
def d(v):
    if v is None: return ''
    if isinstance(v,datetime.datetime): return v.date().isoformat()
    if isinstance(v,(int,float)):
        if 30000<v<60000: return (datetime.date(1899,12,30)+datetime.timedelta(days=int(v))).isoformat()
        return ''
    s=str(v).strip()
    m=re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})$',s)
    if m: return f'{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}'
    m=re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2})$',s)
    if m: return f'20{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}'
    m=re.match(r'^(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})$',s)
    if m and m[2].lower() in BUL: return f'{m[3]}-{BUL[m[2].lower()]:02d}-{int(m[1]):02d}'
    return ''
def n(v):
    if isinstance(v,(int,float)) and not isinstance(v,bool): return round(float(v),2) if v%1 else int(v)
    return 0
def t(v):
    if v is None: return ''
    s=str(v).strip()
    return '' if s in ('-','#DIV/0!') else s
def ck(v): return v is not None and str(v).strip() not in ('','-','x','X')
wb=load_workbook(U+'4a3738a8-2024-11-18_REKAPITULASI_PENJUALAN_UNIT_DAN_REALISASI..xlsx',data_only=True)
ws=wb.worksheets[0]
kpr=[];unparsed=[]
for r in ws.iter_rows(min_row=4,values_only=True):
    r=list(r)+[None]*80
    if r[1] is None and r[2] is None: continue
    no=r[0]
    unit=re.sub(r'\s*-\s*','-',str(r[1] or '').strip())
    for c in (3,4,39,42,64,66):
        if r[c] not in (None,'','-') and not d(r[c]) : unparsed.append((unit,c,r[c]))
    tglAkad=d(r[39]); tempat=t(r[40])
    banks=[t(r[c]) for c in (25,26,27,28) if t(r[c])]
    acc=n(r[13])
    bp=[{'bank':b,'tgl':'','ket':'','hasil':''} for b in banks]
    if bp and (acc>0 or tglAkad):
        key=(tempat.split()[0].upper() if tempat else '')
        idx=next((i for i,x in enumerate(bp) if key and x['bank'].upper().startswith(key)),len(bp)-1)
        bp[idx]['hasil']='ACC'
    if bp and t(r[29]): bp[-1]['ket']=t(r[29])
    pc=[]
    for c in (42,45,48,51,54):
        if n(r[c+1]) or d(r[c]): pc.append({'tgl':d(r[c]),'nominal':n(r[c+1])})
    promo=[x for x in (t(r[68]) and ('Subsidi asuransi: '+t(r[68])),t(r[69]) and ('Bonus: '+t(r[69])),t(r[70]) and ('Subsidi angsuran: '+t(r[70])),t(r[71]) and ('Diskon PPN: '+t(r[71]))) if x]
    rec={'ord':int(no) if isinstance(no,(int,float)) else len(kpr)+1,'unit':unit,'nama':t(r[2]),
      'tglUTJ':d(r[3]),'tglSPR':d(r[4]),'caraBayar':t(r[5]),
      'hargaBank':n(r[6]),'hargaTransaksi':n(r[7]),'utj':n(r[8]),'angsuranUM':n(r[9]),'cashbackUM':n(r[10]),'totalUM':n(r[11]),
      'plafond':n(r[12]),'accBank':acc,'tglACC':'','tum':n(r[14]),
      'berkas':{k:ck(r[c]) for k,c in zip(['ktp','npwp','kk','akta','rk3','suket','slip','rk6','nib','lapkeu'],range(15,25))},
      'bankProses':bp,
      'legal':{k:ck(r[c]) for k,c in zip(['potongPokok','roya','ambilSertifikat','verifikasi','validasi','lunasDP','siLPP','feeKPR','pbg'],range(30,39))},
      'tglAkad':tglAkad,'tempatAkad':tempat,'notaris':t(r[41]),
      'pencairan':pc,'progressBangun':n(r[62]),
      'ajb':ck(r[63]),'tglAJB':d(r[64]),'stu':ck(r[65]),'tglSTU':d(r[66]),
      'keterangan':t(r[67]),'promo':'; '.join(promo)}
    kpr.append(rec)
print('kpr',len(kpr)); print('unparsed',unparsed[:30],len(unparsed))
wb2=load_workbook(U+'34670f74-2026-08-11_DATA_RETENSI_TAHAP_2.xlsx',data_only=True)
ws=wb2.worksheets[0]; ret=[]
for r in ws.iter_rows(min_row=4,values_only=True):
    r=list(r)+[None]*30
    if not r[0] or not r[1]: continue
    ret.append({'ord':len(ret)+1,'blok':t(r[0]),'nama':t(r[1]),'pembayaran':t(r[2]),
     'nilaiUM':n(r[3]),'nilaiKPR':n(r[4]),'terimaUM':n(r[6]),'terimaKPR':n(r[7]),'persenCair':n(r[9]),
     'ret':{'bangunan':round(n(r[10])),'ajb':n(r[11]),'sertifikat':n(r[12]),'pbg':n(r[13]),'pdam':n(r[14]),'listrik':n(r[15]),'pajak':n(r[16])},
     'status':t(r[18]),'bank':t(r[19]),'notaris':t(r[20]),'catatan':'; '.join(x for x in (t(r[21]),t(r[22])) if x)})
print('ret',len(ret)); 
for x in ret[:3]: print(x)
for i,x in enumerate(kpr): json.dump(x,open(f'seed/k{x["ord"]:03d}.json','w'))
for x in ret: json.dump(x,open(f'seed/r{x["ord"]:02d}.json','w'))
import collections
print(collections.Counter(x['caraBayar'] for x in kpr))
print(len(set(x['ord'] for x in kpr)))
print(kpr[-5:][0])
